from openpyxl import Workbook


def generate_batch_report_excel(batch: dict, products: list[dict],
                                stats: dict, file_path: str):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Информация о партии"
    for k, v in batch.items():
        ws1.append([k, v])

    ws2 = wb.create_sheet("Продукция")
    ws2.append(["ID", "Уникальный код", "Аггрегирована", "Дата аггрегации"])
    for p in products:
        ws2.append(
            [
                p["id"],
                p["unique_code"],
                "да" if p["is_aggregated"] else "нет",
                p["aggregated_at"]
            ]
        )

    ws3 = wb.create_sheet("Статистика")
    for k, v in stats.items():
        ws3.append([k, v])

    wb.save(file_path)


def generate_export_batches_excel(batches: list[dict], file_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Партии"

    headers = [
        "Номер партии", "Дата партии", "Статус", "Рабочий центр",
        "Смена", "Бригада", "Номенклатура", "Код ЕКН",
        "Всего продукции", "Аггрегировано"
    ]
    ws.append(headers)

    for b in batches:
        ws.append([
            b["batch_number"], b["batch_date"],
            "Закрыта" if b["is_closed"] else "Открыта",
            b["work_center_id"], b["shift"], b["team"],
            b["nomenclature"], b["ekn_code"],
            b["total_products"], b["aggregated_products"],
        ])

    wb.save(file_path)