from openpyxl import load_workbook

header_lang_map = {
    "СтатусЗакрытия": "is_closed",
    "ПредставлениеЗаданияНаСмену": "task_description",
    "РабочийЦентр": "work_center_name",       # см. замечание ниже про work_center_id
    "Смена": "shift",
    "Бригада": "team",
    "НомерПартии": "batch_number",
    "ДатаПартии": "batch_date",
    "Номенклатура": "nomenclature",
    "КодЕКН": "ekn_code",
    "ИдентификаторРЦ": "work_center_identifier",
    "ДатаВремяНачалаСмены": "shift_start",
    "ДатаВремяОкончанияСмены": "shift_end",
}

def parse_batches_excel(filename):
    wb = load_workbook(filename=filename, read_only=True)
    sheet = wb.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    headers = [header_lang_map[h] for h in headers]

    for row in rows:
        yield dict(zip(headers, row))


    wb.close()